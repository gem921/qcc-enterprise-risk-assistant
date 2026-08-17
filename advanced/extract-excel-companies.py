import json
import re
import sys

from openpyxl import load_workbook


if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")


HEADER_NAMES = {
    "公司名称",
    "企业名称",
    "company",
    "companyname",
    "company_name",
    "compan_name",
    "name",
}


def normalize_header(value):
    text = "" if value is None else str(value)
    text = text.strip().lower()
    text = re.sub(r"[\s\-\u3000]+", "", text)
    return text


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def summarize(values):
    seen = set()
    result = []
    counts = {}
    for value in values:
        text = cell_text(value)
        if not text:
            continue
        counts[text] = counts.get(text, 0) + 1
        if text not in seen:
            seen.add(text)
            result.append(text)
    duplicates = [
        {"name": name, "count": count}
        for name, count in counts.items()
        if count > 1
    ]
    return result, duplicates, sum(counts.values())


def main():
    if len(sys.argv) < 2:
        raise SystemExit("Usage: extract-excel-companies.py <xlsx-file>")

    path = sys.argv[1]
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))

    if not rows:
        print(json.dumps({"companies": []}, ensure_ascii=False))
        return

    header = [normalize_header(value) for value in rows[0]]
    header_names = {normalize_header(value) for value in HEADER_NAMES}
    company_index = next((index for index, name in enumerate(header) if name in header_names), -1)

    if company_index >= 0:
        data_rows = rows[1:]
    else:
        company_index = 0
        data_rows = rows

    raw_values = [row[company_index] if len(row) > company_index else "" for row in data_rows]
    companies, duplicates, raw_count = summarize(raw_values)
    print(json.dumps({
        "companies": companies,
        "rawCount": raw_count,
        "uniqueCount": len(companies),
        "duplicateCount": len(duplicates),
        "duplicates": duplicates,
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
